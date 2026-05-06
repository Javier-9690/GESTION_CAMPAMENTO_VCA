import io
import os
import re
import json
import calendar
from datetime import datetime, date
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import Flask, render_template, request, send_file
from dotenv import load_dotenv

load_dotenv()

SANTIAGO_TZ = ZoneInfo("America/Santiago")


def now_santiago():
    return datetime.now(SANTIAGO_TZ).replace(tzinfo=None)


app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-hotel-key")
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024


@app.context_processor
def inject_now():
    return {"now": now_santiago()}


MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]


HOTEL_HEADERS = [
    "HABITACION",
    "MODULO",
    "RUT",
    "NOMBRE",
    "EMPRESA",
    "N CONTRATO",
    "GERENCIA",
    "SISTEMA TURNO",
    "CALENDARIO ASIGNADO EN SALTO",
    "CO MEL",
    "GENERO",
    "NOMBRE DE TURNO",
]

HOTEL_WIDTHS = [14, 12, 14, 28, 22, 14, 22, 16, 30, 10, 10, 22]


REAL_BEDS = {
    1: 100, 2: 100, 3: 100, 4: 96, 5: 100, 6: 100,
    7: 96, 8: 96, 9: 100, 10: 100, 11: 96, 12: 96,
    13: 24, 14: 24, 15: 24, 16: 64, 17: 64, 18: 64,
    19: 64, 20: 64, 21: 100, 22: 96, 23: 96, 24: 96,
    25: 96, 26: 96, 27: 96, 28: 96, 29: 96, 30: 96,
    31: 96, 32: 96, 33: 96, 34: 144, 35: 144, 36: 144,
    37: 144, 38: 144, 39: 144, 40: 144, 41: 144, 42: 144,
    43: 144, 44: 144, 45: 144, 46: 288, 47: 287,
}


ARAMARK_RED = "B42318"


def thin():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def style_header(cell, fill_hex=ARAMARK_RED):
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin()


def style_cell(cell):
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin()


def clamp_int(value, default, min_value=None, max_value=None):
    try:
        n = int(value)
    except (TypeError, ValueError):
        return default

    if min_value is not None and n < min_value:
        return default

    if max_value is not None and n > max_value:
        return default

    return n


def build_day_status(mes, anio, total_dias):
    """
    Devuelve una lista booleana indicando qué días del reporte ya se deben evaluar.
    Los días futuros no cuentan como vacíos ni como ocupados.
    """
    today = now_santiago().date()
    last_day_of_month = calendar.monthrange(int(anio), int(mes))[1]

    day_is_elapsed = []

    for d in range(1, total_dias + 1):
        if d > last_day_of_month:
            day_is_elapsed.append(False)
            continue

        current_day = date(int(anio), int(mes), d)
        day_is_elapsed.append(current_day <= today)

    total_dias_evaluados = sum(1 for x in day_is_elapsed if x)

    return day_is_elapsed, total_dias_evaluados


def style_report_header(ws):
    fill = PatternFill(start_color=ARAMARK_RED, end_color=ARAMARK_RED, fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    border = thin()

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def autosize_columns(ws, max_width=45):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)

        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))

        ws.column_dimensions[letter].width = min(max_len + 2, max_width)


def add_sheet_from_rows(wb, title, headers, rows):
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)

    for row in rows:
        ws.append(row)

    style_report_header(ws)
    autosize_columns(ws)
    ws.freeze_panes = "A2"

    return ws


def generate_report_excel(payload):
    wb = Workbook()
    wb.remove(wb.active)

    filtros = payload.get("filtros", {})
    kpis = payload.get("kpis", {})
    uso_diario = payload.get("uso_diario", [])
    modulos = payload.get("modulos", [])
    empresas = payload.get("empresas", [])
    habitaciones = payload.get("habitaciones", [])
    habitacion_dia = payload.get("habitacion_dia", [])

    ws = wb.create_sheet("Resumen")
    ws.append(["Campo", "Valor"])
    ws.append(["Módulo", filtros.get("modulo") or "Todos"])
    ws.append(["Empresa", filtros.get("empresa") or "Todas"])
    ws.append(["Turno", filtros.get("turno") or "Todos"])
    ws.append(["Periodo", payload.get("periodo", "")])
    ws.append(["Días calendario", payload.get("total_dias", 0)])
    ws.append(["Días evaluados", payload.get("total_dias_evaluados", 0)])
    ws.append(["Habitaciones filtro", kpis.get("habitaciones", 0)])
    ws.append(["Camas reales filtro", kpis.get("camas_reales", 0)])
    ws.append(["Personas asignadas", kpis.get("asignados", 0)])
    ws.append(["Personas-día", kpis.get("personas_dia", 0)])
    ws.append(["Ocupación habitaciones %", kpis.get("ocupacion_habitaciones", 0)])
    ws.append(["Uso capacidad real %", kpis.get("uso_capacidad_real", 0)])
    style_report_header(ws)
    autosize_columns(ws)

    add_sheet_from_rows(
        wb,
        "Uso diario",
        [
            "Día",
            "Estado",
            "Personas presentes",
            "Habitaciones ocupadas",
            "Camas reales filtro",
            "Ocupación habitaciones %",
            "Uso capacidad real %",
        ],
        [
            [
                r.get("dia"),
                r.get("estado"),
                r.get("personas"),
                r.get("habitaciones_ocupadas"),
                r.get("camas_reales"),
                r.get("ocupacion_habitaciones"),
                r.get("uso_capacidad_real"),
            ]
            for r in uso_diario
        ],
    )

    add_sheet_from_rows(
        wb,
        "Resumen módulos",
        [
            "Módulo",
            "Habitaciones",
            "Camas reales",
            "Asignados",
            "Días ocupados",
            "Personas-día",
            "Ocupación habitaciones %",
            "Uso real %",
        ],
        [
            [
                r.get("modulo"),
                r.get("habitaciones"),
                r.get("camas_reales"),
                r.get("asignados"),
                r.get("dias_ocupados"),
                r.get("personas_dia"),
                r.get("ocupacion_habitaciones"),
                r.get("uso_real"),
            ]
            for r in modulos
        ],
    )

    add_sheet_from_rows(
        wb,
        "Resumen empresas",
        [
            "Empresa",
            "Habitaciones",
            "Asignados",
            "Días perdidos",
            "Ocupación %",
        ],
        [
            [
                r.get("empresa"),
                r.get("habitaciones"),
                r.get("asignados"),
                r.get("dias_perdidos"),
                r.get("ocupacion"),
            ]
            for r in empresas
        ],
    )

    add_sheet_from_rows(
        wb,
        "Detalle habitaciones",
        [
            "Módulo",
            "Habitación",
            "Camas reales",
            "Días ocupados",
            "Días vacíos",
            "Ocupación %",
            "Empresas",
            "Turnos",
            "Asignados",
        ],
        [
            [
                r.get("modulo"),
                r.get("habitacion"),
                r.get("camas_reales"),
                r.get("dias_ocupados"),
                r.get("dias_vacios"),
                r.get("ocupacion"),
                r.get("empresas"),
                r.get("turnos"),
                r.get("asignados"),
            ]
            for r in habitaciones
        ],
    )

    if habitacion_dia:
        max_days = max(len(r.get("dias", [])) for r in habitacion_dia)
    else:
        max_days = 0

    day_headers = [f"Día {i}" for i in range(1, max_days + 1)]

    add_sheet_from_rows(
        wb,
        "Habitación día",
        ["Módulo", "Habitación", "Camas reales"] + day_headers,
        [
            [
                r.get("modulo"),
                r.get("habitacion"),
                r.get("camas_reales"),
            ] + r.get("dias", [])
            for r in habitacion_dia
        ],
    )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def generate_hoteleria():
    wb = Workbook()
    ws = wb.active
    ws.title = "Base Hoteleria"
    ws.row_dimensions[1].height = 36

    for col, h in enumerate(HOTEL_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        style_header(cell)

    for i, w in enumerate(HOTEL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:" + get_column_letter(len(HOTEL_HEADERS)) + "1"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def generate_censo(mes, anio):
    days = calendar.monthrange(int(anio), int(mes))[1]

    wb = Workbook()
    ws = wb.active
    ws.title = "Censo Ontracking"
    ws.row_dimensions[1].height = 30

    headers = ["Modulo", "Habitacion"] + [str(d) for d in range(1, days + 1)]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        style_header(cell, "8B0000")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16

    for d in range(1, days + 1):
        ws.column_dimensions[get_column_letter(d + 2)].width = 5

    ws.freeze_panes = "C2"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def read_rows(file_obj):
    wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def normalize(val):
    if val is None:
        return ""

    s = str(val).strip()

    if re.match(r"^\d+\.0+$", s):
        s = str(int(float(s)))

    return s


def extract_mod_num(mod_str):
    s = normalize(mod_str).upper()

    s = re.sub(r"^M[OÓ]DULO\s*", "", s)
    s = re.sub(r"^MOD\s*", "", s)
    s = re.sub(r"^M\s*", "", s)

    m = re.match(r"^0*(\d+)", s)

    return int(m.group(1)) if m else None


def get_real_beds(mod_str):
    n = extract_mod_num(mod_str)
    return REAL_BEDS.get(n, None)


def normalize_suffix(s):
    s = normalize(s).strip().upper()

    m = re.match(r"^0*(\d+)([A-Z]*)$", s)

    if m:
        return str(int(m.group(1))) + m.group(2)

    return s


def census_key(mod_val, hab_val):
    mod_num = extract_mod_num(mod_val)
    hab = normalize(hab_val).upper()

    m = re.match(r"^M\d+H(.+)$", hab)

    if m:
        suffix = normalize_suffix(m.group(1))
    else:
        suffix = normalize_suffix(hab) if hab else hab

    return mod_num, suffix


def hotel_key(mod_val, hab_val):
    mod_num = extract_mod_num(mod_val)
    hab = normalize(hab_val).upper()

    if mod_num is not None:
        prefix = str(mod_num)

        if hab.startswith(prefix):
            suffix = normalize_suffix(hab[len(prefix):])
        else:
            suffix = normalize_suffix(hab)
    else:
        suffix = normalize_suffix(hab) if hab else hab

    return mod_num, suffix


def process(hotel_file, censo_file, mes_reporte=None, anio_reporte=None):
    hotel_rows = read_rows(hotel_file)

    if not hotel_rows:
        raise ValueError("Plantilla hotelería vacía")

    hotel_headers = [normalize(c).upper() for c in hotel_rows[0]]

    def find_col(fragments):
        for i, h in enumerate(hotel_headers):
            if any(f in h for f in fragments):
                return i
        return None

    idx_hab = find_col(["HABITACI"])
    idx_mod = find_col(["MODULO", "MÓDULO"])
    idx_nom = find_col(["NOMBRE"])
    idx_rut = find_col(["RUT"])
    idx_emp = find_col(["EMPRESA"])
    idx_ger = find_col(["GERENCIA"])
    idx_gen = find_col(["GENERO", "GÉNERO"])
    idx_con = find_col(["CONTRATO"])
    idx_cal = find_col(["CALENDARIO", "SALTO"])
    idx_mel = find_col(["MEL"])

    idx_sis_tur = find_col(["SISTEMA TURNO"])
    idx_nom_tur = find_col(["NOMBRE DE TURNO"])

    if idx_hab is None or idx_mod is None:
        raise ValueError(
            "Plantilla hotelería: faltan columnas HABITACION/MODULO. "
            f"Columnas detectadas: {hotel_headers}"
        )

    hotel_data = {}
    hotel_rows_loaded = 0

    for row in hotel_rows[1:]:
        if not any(row):
            continue

        raw_hab = normalize(row[idx_hab] if idx_hab < len(row) else "")
        raw_mod = normalize(row[idx_mod] if idx_mod < len(row) else "")

        if not raw_hab or not raw_mod:
            continue

        def get(idx):
            return normalize(row[idx] if idx is not None and idx < len(row) else "")

        persona = {
            "nombre": get(idx_nom),
            "rut": get(idx_rut),
            "empresa": get(idx_emp),
            "gerencia": get(idx_ger),
            "sistema_turno": get(idx_sis_tur),
            "turno": get(idx_nom_tur) or get(idx_sis_tur),
            "genero": get(idx_gen),
            "contrato": get(idx_con),
            "calendario": get(idx_cal),
            "co_mel": get(idx_mel),
        }

        key = hotel_key(raw_mod, raw_hab)
        hotel_data.setdefault(key, []).append(persona)
        hotel_rows_loaded += 1

    censo_rows = read_rows(censo_file)

    if not censo_rows:
        raise ValueError("Plantilla censo vacía")

    censo_headers = [normalize(c) for c in censo_rows[0]]
    day_indices = [(i, h) for i, h in enumerate(censo_headers) if h.isdigit()]
    total_dias = len(day_indices)

    if total_dias == 0:
        raise ValueError("Plantilla censo: no se encontraron columnas de días 1, 2, 3...")

    now = now_santiago()

    mes_reporte = clamp_int(mes_reporte, now.month, 1, 12)
    anio_reporte = clamp_int(anio_reporte, now.year, now.year - 10, now.year + 10)

    day_is_elapsed, total_dias_evaluados = build_day_status(
        mes_reporte,
        anio_reporte,
        total_dias,
    )

    mod_col2 = 0
    hab_col2 = 1

    for i, h in enumerate(censo_headers):
        hu = h.upper()

        if "MODULO" in hu or "MÓDULO" in hu:
            mod_col2 = i
        elif "HABITACI" in hu:
            hab_col2 = i

    rooms = []
    matched_count = 0

    for row in censo_rows[1:]:
        if not any(row):
            continue

        raw_mod = normalize(row[mod_col2] if mod_col2 < len(row) else "")
        raw_hab = normalize(row[hab_col2] if hab_col2 < len(row) else "")

        if not raw_mod or not raw_hab:
            continue

        key = census_key(raw_mod, raw_hab)

        daily = []

        for i, _ in day_indices:
            v = row[i] if i < len(row) else None

            try:
                daily.append(int(float(str(v))) if v not in (None, "", "None") else 0)
            except (ValueError, TypeError):
                daily.append(0)

        daily_evaluado = [
            value if idx < len(day_is_elapsed) and day_is_elapsed[idx] else 0
            for idx, value in enumerate(daily)
        ]

        dias_ocupados = sum(1 for v in daily_evaluado if v > 0)
        personas_dia = sum(daily_evaluado)

        asignados = hotel_data.get(key, [])

        if asignados:
            matched_count += 1

        empresas = sorted({p["empresa"] for p in asignados if p["empresa"]})
        turnos = sorted({p["turno"] for p in asignados if p["turno"]})

        rooms.append({
            "modulo": raw_mod,
            "mod_num": key[0],
            "habitacion": raw_hab,
            "key_str": f"{key[0]}-{key[1]}",
            "daily": daily,
            "dias_ocupados": dias_ocupados,
            "dias_vacios": max(total_dias_evaluados - dias_ocupados, 0),
            "personas_dia": personas_dia,
            "pct_ocupacion": round(dias_ocupados / total_dias_evaluados * 100, 1) if total_dias_evaluados else 0,
            "asignados": asignados,
            "n_asignados": len(asignados),
            "empresas": empresas,
            "turnos": turnos,
        })

    if not rooms:
        raise ValueError("No se encontraron habitaciones en el censo")

    rooms_by_module = {}

    for room in rooms:
        rooms_by_module.setdefault(room["modulo"], []).append(room)

    for modulo, module_rooms in rooms_by_module.items():
        camas_reales_modulo = get_real_beds(modulo) or 0
        habitaciones_modulo = len(module_rooms)

        camas_reales_por_habitacion = (
            camas_reales_modulo / habitaciones_modulo
            if habitaciones_modulo and camas_reales_modulo
            else 0
        )

        for room in module_rooms:
            room["camas_reales"] = round(camas_reales_por_habitacion, 2)

    hotel_sample = [
        {"mod": str(k[0]), "hab": str(k[1])}
        for k in sorted(hotel_data.keys())[:5]
    ]

    censo_sample = []

    for row in censo_rows[1:]:
        if not any(row):
            continue

        rm = normalize(row[mod_col2] if mod_col2 < len(row) else "")
        rh = normalize(row[hab_col2] if hab_col2 < len(row) else "")

        if rm and rh:
            k = census_key(rm, rh)
            censo_sample.append({"mod": str(k[0]), "hab": str(k[1])})

        if len(censo_sample) >= 5:
            break

    match_info = {
        "hotel_rows": hotel_rows_loaded,
        "censo_rooms": len(rooms),
        "matched": matched_count,
        "hotel_sample": hotel_sample,
        "censo_sample": censo_sample,
    }

    modulos = {}

    for r in rooms:
        m = r["modulo"]

        if m not in modulos:
            camas_reales = get_real_beds(m)

            modulos[m] = {
                "modulo": m,
                "mod_num": extract_mod_num(m),
                "habitaciones": 0,
                "camas_reales": camas_reales,
                "personas_asignadas": 0,
                "dias_ocupados": 0,
                "dias_totales": 0,
                "personas_dia_total": 0,
            }

        modulos[m]["habitaciones"] += 1
        modulos[m]["personas_asignadas"] += r["n_asignados"]
        modulos[m]["dias_ocupados"] += r["dias_ocupados"]
        modulos[m]["dias_totales"] += total_dias_evaluados
        modulos[m]["personas_dia_total"] += r["personas_dia"]

    for m in modulos.values():
        m["pct_ocupacion"] = round(
            m["dias_ocupados"] / m["dias_totales"] * 100,
            1
        ) if m["dias_totales"] else 0

        cr = m["camas_reales"]

        if cr and total_dias_evaluados:
            m["pct_cap_real"] = round(
                m["personas_dia_total"] / (cr * total_dias_evaluados) * 100,
                1
            )
        else:
            m["pct_cap_real"] = None

    module_list = sorted(
        modulos.values(),
        key=lambda x: x["mod_num"] or 9999
    )

    empresas = {}

    def get_emp(name):
        if name not in empresas:
            empresas[name] = {
                "empresa": name,
                "habitaciones": 0,
                "personas": 0,
                "dias_ocupados": 0,
                "dias_totales": 0,
                "dias_perdidos": 0,
            }

        return empresas[name]

    for r in rooms:
        for p in r["asignados"]:
            get_emp(p["empresa"] or "Sin empresa")["personas"] += 1

        tag_emps = r["empresas"] if r["empresas"] else (
            ["Sin asignar"] if not r["asignados"] else []
        )

        for emp in tag_emps:
            e = get_emp(emp)
            e["habitaciones"] += 1
            e["dias_ocupados"] += r["dias_ocupados"]
            e["dias_totales"] += total_dias_evaluados
            e["dias_perdidos"] += r["dias_vacios"]

    for e in empresas.values():
        e["pct_ocupacion"] = round(
            e["dias_ocupados"] / e["dias_totales"] * 100,
            1
        ) if e["dias_totales"] else 0

    empresa_list = sorted(
        empresas.values(),
        key=lambda x: x["dias_perdidos"],
        reverse=True
    )

    total_hab = len(rooms)
    total_asig = sum(r["n_asignados"] for r in rooms)
    total_oc = sum(r["dias_ocupados"] for r in rooms)
    total_personas_dia = sum(r["personas_dia"] for r in rooms)

    total_pct = round(
        total_oc / (total_hab * total_dias_evaluados) * 100,
        1
    ) if total_hab * total_dias_evaluados else 0

    total_camas_reales = sum(
        m["camas_reales"] or 0
        for m in modulos.values()
    )

    total_pct_cap_real = round(
        total_personas_dia / (total_camas_reales * total_dias_evaluados) * 100,
        1
    ) if total_camas_reales and total_dias_evaluados else 0

    company_names = sorted({
        p["empresa"]
        for r in rooms
        for p in r["asignados"]
        if p["empresa"]
    })

    turno_names = sorted({
        p["turno"]
        for r in rooms
        for p in r["asignados"]
        if p["turno"]
    })

    modulo_names = sorted(
        {r["modulo"] for r in rooms},
        key=lambda x: extract_mod_num(x) or 9999
    )

    sorted_rooms = sorted(
        rooms,
        key=lambda x: (
            extract_mod_num(x["modulo"]) or 9999,
            normalize(x["habitacion"])
        )
    )

    module_capacity_map = {
        m["modulo"]: m["camas_reales"] or 0
        for m in module_list
    }

    rooms_js = [
        {
            "modulo": r["modulo"],
            "mod_num": r["mod_num"],
            "habitacion": r["habitacion"],
            "empresas": r["empresas"],
            "turnos": r["turnos"],
            "daily": r["daily"],
            "dias_ocupados": r["dias_ocupados"],
            "dias_vacios": r["dias_vacios"],
            "personas_dia": r["personas_dia"],
            "camas_reales": r.get("camas_reales", 0),
            "pct_ocupacion": r["pct_ocupacion"],
            "n_asignados": r["n_asignados"],
            "asignados": [
                {
                    "empresa": p.get("empresa", ""),
                    "turno": p.get("turno", ""),
                    "nombre": p.get("nombre", ""),
                    "rut": p.get("rut", ""),
                }
                for p in r["asignados"]
            ],
        }
        for r in sorted_rooms
    ]

    return {
        "total_habitaciones": total_hab,
        "total_asignadas": total_asig,
        "total_camas_reales": total_camas_reales,
        "total_personas_dia": total_personas_dia,
        "total_pct_cap_real": total_pct_cap_real,
        "total_dias": total_dias,
        "total_dias_evaluados": total_dias_evaluados,
        "day_is_elapsed": day_is_elapsed,
        "mes_reporte": mes_reporte,
        "anio_reporte": anio_reporte,
        "periodo": f"{MESES[mes_reporte - 1]} {anio_reporte}",
        "pct_ocupacion": total_pct,
        "modulos": module_list,
        "empresas": empresa_list,
        "rooms": sorted_rooms,
        "rooms_js": rooms_js,
        "company_names": company_names,
        "turno_names": turno_names,
        "modulo_names": modulo_names,
        "module_capacity_map": module_capacity_map,
        "match_info": match_info,
        "day_labels": list(range(1, total_dias + 1)),
    }


@app.route("/")
def index():
    now = now_santiago()

    return render_template(
        "index.html",
        mes=now.month,
        anio=now.year,
        meses=MESES,
        anios=list(range(now.year - 1, now.year + 2)),
    )


@app.route("/download/hoteleria")
def download_hoteleria():
    out = generate_hoteleria()

    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="plantilla_hoteleria.xlsx",
    )


@app.route("/download/censo")
def download_censo():
    now = now_santiago()

    mes = request.args.get("mes", now.month, type=int)
    anio = request.args.get("anio", now.year, type=int)

    if mes < 1 or mes > 12:
        mes = now.month

    if anio < now.year - 5 or anio > now.year + 5:
        anio = now.year

    out = generate_censo(mes, anio)

    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"plantilla_censo_{anio}_{str(mes).zfill(2)}.xlsx",
    )


@app.route("/procesar", methods=["POST"])
def procesar():
    hotel_file = request.files.get("hoteleria")
    censo_file = request.files.get("censo")
    now = now_santiago()

    if not hotel_file or not censo_file:
        return render_template(
            "index.html",
            error="Debes subir ambas plantillas.",
            mes=now.month,
            anio=now.year,
            meses=MESES,
            anios=list(range(now.year - 1, now.year + 2)),
        )

    try:
        mes_reporte = request.form.get("mes_reporte", type=int)
        anio_reporte = request.form.get("anio_reporte", type=int)

        results = process(hotel_file, censo_file, mes_reporte, anio_reporte)
        return render_template("resultados.html", r=results)

    except Exception as e:
        return render_template(
            "index.html",
            error=f"Error al procesar: {str(e)}",
            mes=now.month,
            anio=now.year,
            meses=MESES,
            anios=list(range(now.year - 1, now.year + 2)),
        )


@app.route("/download/reporte", methods=["POST"])
def download_reporte():
    raw = request.form.get("report_json", "")

    if not raw:
        return "No se recibió información para generar el reporte.", 400

    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        return "El reporte recibido no tiene formato válido.", 400

    out = generate_report_excel(payload)

    periodo = payload.get("periodo", "reporte").replace(" ", "_").lower()

    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"reporte_hotel_analytics_{periodo}.xlsx",
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5001)
